"""Convert the supplied CiteScore workbook and report annexes to site data."""

from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import re
from collections import OrderedDict
from datetime import date
from pathlib import Path

from journal_data import BEHAVIOUR_FILE, CITESCORE_FILE, normalize_journal_name


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def clean_issn(value: object) -> str:
    if value in (None, ""):
        return ""
    text = re.sub(r"[^0-9X]", "", str(value).upper())
    if len(text) < 8:
        text = text.zfill(8)
    return f"{text[:4]}-{text[4:8]}" if len(text) == 8 else text


def optional_number(value: object) -> int | float | None:
    if value in (None, ""):
        return None
    number = float(value)
    return int(number) if number.is_integer() else number


def import_citescore(workbook_path: Path) -> dict:
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    columns = {name: index for index, name in enumerate(next(rows))}
    journals: OrderedDict[str, dict] = OrderedDict()
    for row in rows:
        source_id = str(row[columns["Scopus Source ID"]] or "")
        title = str(row[columns["Title"]] or "").strip()
        if not source_id or not title:
            continue
        journal = journals.setdefault(
            source_id,
            {
                "source_id": source_id,
                "title": title,
                "issns": [],
                "citescore": optional_number(row[columns["CiteScore"]]),
                "snip": optional_number(row[columns["SNIP"]]),
                "sjr": optional_number(row[columns["SJR"]]),
                "open_access": str(row[columns["Open Access"]] or "").upper() == "YES",
                "url": str(row[columns["URL Scopus Source ID"]] or ""),
                "categories": [],
            },
        )
        for field in ("Print ISSN", "E-ISSN"):
            issn = clean_issn(row[columns[field]])
            if issn and issn not in journal["issns"]:
                journal["issns"].append(issn)
        quartile = optional_number(row[columns["Quartile"]])
        journal["categories"].append(
            {
                "code": str(row[columns["Scopus ASJC Code (Sub-subject Area)"]] or ""),
                "name": str(row[columns["Scopus Sub-Subject Area"]] or ""),
                "percentile": optional_number(row[columns["Percentile"]]),
                "rank": optional_number(row[columns["RANK"]]),
                "rank_out_of": optional_number(row[columns["Rank Out Of"]]),
                "quartile": f"Q{quartile}" if quartile else "",
                "top_10_percent": bool(row[columns["Top 10% (CiteScore Percentile)"]]),
            }
        )
    for journal in journals.values():
        quartiles = [category["quartile"] for category in journal["categories"] if category["quartile"]]
        journal["best_quartile"] = min(quartiles, key=lambda value: int(value[1:])) if quartiles else ""
        journal["top_10_percent"] = any(category["top_10_percent"] for category in journal["categories"])
    return {
        "metadata": {
            "source": workbook_path.name,
            "database": "Scopus CiteScore 2025 annual values",
            "ranking_system": "CiteScore",
            "not_jcr": True,
            "source_sha256": sha256(workbook_path),
            "imported_on": date.today().isoformat(),
            "journal_count": len(journals),
            "note": "This dataset is Scopus CiteScore, not Clarivate Journal Citation Reports (JCR).",
        },
        "journals": list(journals.values()),
    }


def bullet_titles(lines: list[str], start: int, end: int) -> list[str]:
    return [line[2:].strip() for line in lines[start:end] if line.startswith("- ")]


def import_behaviour_report(report_path: Path) -> dict:
    lines = report_path.read_text(encoding="utf-8").splitlines()
    markers = {}
    for index, line in enumerate(lines):
        if line.startswith("## Anexo II"):
            markers["annex_2"] = index
        elif line.startswith("## Anexo I"):
            markers["annex_1"] = index
        elif line.startswith("### Nivel: Muy alto"):
            markers["very_high"] = index
        elif line.startswith("### Nivel: Alto"):
            markers["high"] = index
        elif line.startswith("### Nivel: Moderado"):
            markers["moderate"] = index
    missing = {"annex_1", "annex_2", "very_high", "high", "moderate"} - markers.keys()
    if missing:
        raise RuntimeError(f"Missing report sections: {sorted(missing)}")
    annex_1 = bullet_titles(lines, markers["annex_1"], markers["annex_2"])
    very_high = bullet_titles(lines, markers["very_high"], markers["high"])
    high = bullet_titles(lines, markers["high"], markers["moderate"])
    moderate = bullet_titles(lines, markers["moderate"], len(lines))
    records: dict[str, dict] = {}
    for title in annex_1:
        records[normalize_journal_name(title)] = {
            "title": title,
            "access_model": "open_access",
            "level": "non_standard",
            "level_confidence": "report_annex",
        }
    for title, level, confidence in (
        *((title, "very_high", "high") for title in very_high),
        *((title, "high", "inferred_from_transcription") for title in high),
        *((title, "moderate", "inferred_from_transcription") for title in moderate),
    ):
        records[normalize_journal_name(title)] = {
            "title": title,
            "access_model": "subscription",
            "level": level,
            "level_confidence": confidence,
        }
    return {
        "metadata": {
            "source": report_path.name,
            "original_document": "210930_Openaccess_Revistas_baneadas_.pdf",
            "authors": [
                "M. Ángeles Oviedo García",
                "José Carlos Casillas Bueno",
                "M. Rosario González Rodríguez",
            ],
            "published": "2021-09-30",
            "study_period": "2017-2019",
            "source_sha256": sha256(report_path),
            "imported_on": date.today().isoformat(),
            "journal_count": len(records),
            "note": (
                "The report identifies non-standard bibliometric behaviour in self-citation and/or citable-item volume. "
                "It is not an official ANECA blacklist, ban, or journal endorsement list."
            ),
        },
        "journals": sorted(records.values(), key=lambda item: normalize_journal_name(item["title"])),
    }


def write_gzip_json(path: Path, payload: dict) -> None:
    encoded = (json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n").encode("utf-8")
    with path.open("wb") as raw:
        with gzip.GzipFile(filename="", mode="wb", fileobj=raw, mtime=0) as compressed:
            compressed.write(encoded)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("report", type=Path)
    args = parser.parse_args()
    CITESCORE_FILE.parent.mkdir(parents=True, exist_ok=True)
    citescore = import_citescore(args.workbook)
    behaviour = import_behaviour_report(args.report)
    write_gzip_json(CITESCORE_FILE, citescore)
    BEHAVIOUR_FILE.write_text(json.dumps(behaviour, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Imported {len(citescore['journals'])} CiteScore journals and {len(behaviour['journals'])} report journals")


if __name__ == "__main__":
    main()
